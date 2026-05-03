import hashlib
import json
import logging
from pathlib import Path

import docx as _docx
import pypdf
from langchain_core.documents import Document

from local_agent.config import Settings

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None


logger = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def _load_txt(file_path: str) -> str:
    encodings_to_try = ["utf-8", "cp1252", "latin-1", "ascii"]
    for encoding in encodings_to_try:
        try:
            with open(file_path, encoding=encoding) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise RuntimeError(f"Could not decode {file_path} with any of {encodings_to_try}")


def _load_pdf(file_path: str, settings: Settings) -> str:
    path = Path(file_path)
    cache_dir = Path(settings.pdf_cache_dir)

    stat = path.stat()
    cache_key_raw = (
        f"{path.resolve()}|{stat.st_size}|{stat.st_mtime_ns}|"
        f"{settings.pdf_max_pages}|{settings.pdf_extractor}"
    )
    cache_key = hashlib.sha256(cache_key_raw.encode("utf-8")).hexdigest()
    cache_file = cache_dir / f"{cache_key}.txt"

    if cache_file.exists():
        return cache_file.read_text(encoding="utf-8")

    if settings.pdf_extractor in {"auto", "pymupdf"} and fitz is not None:
        with fitz.open(file_path) as doc:
            page_count = len(doc)
            limit = page_count if settings.pdf_max_pages <= 0 else min(settings.pdf_max_pages, page_count)
            pages = [doc[i].get_text("text") or "" for i in range(limit)]
            text = "\n".join(pages)
    else:
        reader = pypdf.PdfReader(file_path)
        page_count = len(reader.pages)
        limit = page_count if settings.pdf_max_pages <= 0 else min(settings.pdf_max_pages, page_count)
        pages = [reader.pages[i].extract_text() or "" for i in range(limit)]
        text = "\n".join(pages)

    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(text, encoding="utf-8")
    return text


def _load_docx(file_path: str) -> str:
    doc = _docx.Document(file_path)
    return "\n".join(para.text for para in doc.paragraphs)


def load_documents(settings: Settings) -> list[Document]:
    docs_path = Path(settings.docs_dir)
    if not docs_path.exists():
        raise FileNotFoundError(f"Docs directory not found: {docs_path.resolve()}")

    source_files = sorted(
        p for p in docs_path.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )
    logger.info(
        "Scanning %d supported files (.txt/.pdf/.docx) in %s",
        len(source_files),
        docs_path.resolve(),
    )

    docs: list[Document] = []
    for index, path in enumerate(source_files, start=1):
        file_path = str(path)
        ext = path.suffix.lower()
        try:
            if ext == ".txt":
                text = _load_txt(file_path)
            elif ext == ".pdf":
                text = _load_pdf(file_path, settings)
            elif ext == ".docx":
                text = _load_docx(file_path)
            else:
                continue

            docs.append(Document(page_content=text, metadata={"source": file_path, "type": ext.lstrip(".")}))
            logger.info("Loaded [%d/%d] %s (%s)", index, len(source_files), path.name, ext)
        except Exception as exc:
            raise RuntimeError(f"Failed to load {file_path}: {exc}") from exc

    if not docs:
        exts = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"No supported files ({exts}) found in {docs_path.resolve()}")

    logger.info("Loaded %d documents successfully", len(docs))
    return docs


def collect_docs_state(docs_dir: str) -> dict:
    docs_path = Path(docs_dir)
    source_files = sorted(
        p for p in docs_path.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    files: dict[str, dict[str, int]] = {}
    for path in source_files:
        stat = path.stat()
        files[str(path.resolve())] = {"size": int(stat.st_size), "mtime_ns": int(stat.st_mtime_ns)}

    return {"docs_dir": str(docs_path.resolve()), "file_count": len(files), "files": files}


def load_saved_docs_state(state_file: str) -> dict | None:
    path = Path(state_file)
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        logger.warning("Failed to read docs state file, forcing rebuild: %s", state_file)
        return None


def save_docs_state(state_file: str, state: dict) -> None:
    Path(state_file).write_text(json.dumps(state, indent=2, sort_keys=True), encoding="utf-8")


def docs_changed(docs_dir: str, state_file: str) -> bool:
    current = collect_docs_state(docs_dir)
    saved = load_saved_docs_state(state_file)
    return saved != current

